"""FDRA thresholds workbook (site/thresholds/FDRA_thresholds.xlsx) <-> fdop_start.json.

The workbook is the place to change FDOP thresholds. Edit it in Excel, upload it to GitHub, and
the "Thresholds check" workflow reports any problems; the site build then uses it. Rows with errors
are skipped and that FDRA keeps its last good values.

  python site/thresholds_book.py export [json]      # write the workbook from fdop_start.json
  python site/thresholds_book.py check              # validate; exit 1 on errors (CI)
  python site/thresholds_book.py build OUT [prev]   # write OUT/fdop_start.json (+ OUT/thresholds_check.json)
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / 'site' / 'thresholds' / 'FDRA_thresholds.xlsx'
FDRAS = ROOT / 'docs' / 'data' / 'fdras.geojson'
LEGACY = ROOT / 'docs' / 'data' / 'fdop_start.json'

INDEX = {'ERC': 'erc', 'BI': 'bi', 'IC': 'ic', 'SC': 'sc', 'KBDI': 'kbdi', 'GSI': 'gsi', 'GSI (inverted)': 'gsic',
         '100-hr FM': 'f100', '1000-hr FM': 'f1000'}
INDEX_NAME = {v: k for k, v in INDEX.items()}
FUEL = ['V', 'W', 'X', 'Y', 'Z']
NB, NC = 6, 7
ABOUT = ('Starting points for the FDOP Workbench and the popup levels, built from site/thresholds/FDRA_thresholds.xlsx. '
         'Breakpoint = top of its class. Optional "weights": {stationId: weight} gives a weighted SIG; stations without a weight count 1.')

INFO_COLS = ['FDRA ID', 'FDRA name (info only)', 'Agency (info only)', 'FDOP source', 'Primary index', 'Fuel model',
             'SIG stations', 'Station weights', 'Season start (MM-DD)', 'Season end (MM-DD)', 'Calibration from', 'Calibration to',
             'Large fire (acres)', 'Multi-fire day (fires)', 'Class names (primary)', 'Advanced (JSON)', 'Notes']
TOOL_COLS = ['FDRA ID', 'Order', 'Tool name', 'Index', 'Fuel model'] + [f'Break {i}' for i in range(1, NB + 1)] + \
            [f'Class {i}' for i in range(1, NC + 1)]
TOOL_OPT = ['Days used', 'Adjustments']          # optional columns (older workbooks without them still read)
MAT_COLS = (['FDRA ID', 'Order', 'Tool name', 'Row index', 'Row fuel model'] + [f'Row break {i}' for i in range(1, NB + 1)] +
            ['Column index', 'Column fuel model'] + [f'Column break {i}' for i in range(1, NB + 1)] +
            [f'Row {i} levels' for i in range(1, NC + 1)] + [f'Class {i}' for i in range(1, NC + 1)] + TOOL_OPT)
ADJ_VARS = {'TMAX': 'tmax', 'RH MIN': 'rhmin', '100-HR FM': 'f100', '1000-HR FM': 'f1000', 'GSI': 'gsi', 'KBDI': 'kbdi',
            'ERC': 'erc', 'BI': 'bi', 'IC': 'ic', 'SC': 'sc'}
ADJ_NAME = {'tmax': 'Tmax', 'rhmin': 'RH min', 'f100': '100-hr FM', 'f1000': '1000-hr FM', 'gsi': 'GSI', 'kbdi': 'KBDI',
            'erc': 'ERC', 'bi': 'BI', 'ic': 'IC', 'sc': 'SC'}

HOWTO = """HOW TO UPDATE THE FDRA THRESHOLDS
1. Download this file from GitHub: site/thresholds/FDRA_thresholds.xlsx, then Download raw file.
2. Edit it in Excel. Keep the sheet names and the column headings as they are.
   - Thresholds sheet: one row per decision tool (staffing, adjective, dispatch, and so on).
     Breakpoints go in Break 1, Break 2, and so on, as numbers. A breakpoint is the TOP of its class:
     a value above Break 1 is in class 2. Class names go in Class 1, Class 2, and so on, and there
     must be one more class than there are breakpoints.
     Use Index and Fuel model from the dropdowns. Order sets the order the tools are listed in.
     The first tool for an index sets that index's default breakpoints in the Workbench.
     Days used (optional): which days the tool averages. Blank = the observed day and the forecast day.
     F1 = the next forecast day only; O1+F1+F2+F3 = average of yesterday's observed and the next three forecast days.
     Adjustments (optional): move the class up or down when a condition is met, one per line or separated by ;
     for example  Tmax < 57 : -1   or   GSI > 0.3 : -1   or   100-hr FM > 14 : -1 @F1  (the @ part sets its own days).
     Variables: Tmax, RH min, 100-hr FM, 1000-hr FM, GSI, KBDI, ERC, BI, IC, SC.
   - Matrix tools sheet: a level read from a table, such as BI rows x ERC columns. Give the row index and
     breakpoints, the column index and breakpoints, then one "Row N levels" cell per row with the level for each
     column, separated by spaces (for example 1 1 1 2 3). Levels are numbers from 1 to the number of classes.
   - FDOP info sheet: one row per FDRA (the plan source, SIG stations, season and so on).
     Rows with no FDOP source and no tools are ignored.
   - To change a threshold, edit its row. To add a tool, add a row with the FDRA ID.
     To remove a tool, delete the row.
3. Upload it to GitHub: open the site/thresholds folder, then Add file, then Upload files. Drop the
   file in, write a short note (for example "LPF 2027 FDOP"), and Commit changes.
4. The "Thresholds check" run starts on its own. A green check means the file is good. A red X lists
   each problem by sheet and row. Fix those rows and upload again. FDRAs with errors keep their last
   good values; everything else updates.
5. The site and the AGOL popups pick up the new thresholds at the next site build, which starts on
   its own after the upload (about 15-30 minutes).
"""


# ---------------------------------------------------------------- export
def export(start: dict, path: Path = BOOK):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    fd = json.loads(FDRAS.read_text())['features']
    meta = {f['properties']['fdra_id']: f['properties'] for f in fd}
    wb = Workbook()
    ws = wb.active; ws.title = 'How to update'
    for i, line in enumerate(HOWTO.splitlines(), 1):
        ws.cell(i, 1, line).font = Font(bold=(i == 1))
    ws.column_dimensions['A'].width = 120
    hdr = PatternFill('solid', fgColor='DDE6F0')

    def sheet(name, cols, widths):
        s = wb.create_sheet(name)
        s.append(cols)
        for c, w in zip(s[1], widths):
            c.font = Font(bold=True); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='top')
            s.column_dimensions[c.column_letter].width = w
        s.freeze_panes = 'B2'
        return s

    tl = sheet('Thresholds', TOOL_COLS + TOOL_OPT, [22, 7, 48, 13, 9] + [8] * NB + [14] * NC + [16, 34])
    mt = sheet('Matrix tools', MAT_COLS, [22, 7, 48, 13, 9] + [8] * NB + [13, 9] + [8] * NB + [12] * NC + [14] * NC + [16, 34])
    info = sheet('FDOP info', INFO_COLS, [22, 26, 18, 50, 13, 9, 26, 22, 11, 11, 10, 10, 10, 10, 30, 24, 80])
    ids = sorted(set(meta) | set(start))
    for k in ids:
        e = start.get(k, {}); m = meta.get(k, {})
        extra = {x: e[x] for x in e if x not in ('source', 'stations', 'fm', 'idx', 'notes', 'breaks', 'classes', 'tools',
                                                 'season', 'calib', 'large', 'mfd', 'weights', 'matrix_tools')}
        info.append([k, m.get('fdra_name', ''), m.get('agency', ''), e.get('source', ''), INDEX_NAME.get(e.get('idx'), ''),
                     e.get('fm', ''), ', '.join(map(str, e.get('stations', []))),
                     ', '.join(f'{a}:{b}' for a, b in (e.get('weights') or {}).items()),
                     *(e.get('season') or ['', '']), *(e.get('calib') or ['', '']), e.get('large', ''), e.get('mfd', ''),
                     ' | '.join(e.get('classes') or []), json.dumps(extra) if extra else '', e.get('notes', '')])
        for i, t in enumerate(e.get('tools', []), 1):
            tl.append([k, t.get('order', i), t['label'], INDEX_NAME[t['idx']], t['fm']] + t['breaks'] + [None] * (NB - len(t['breaks'])) +
                      t['classes'] + [None] * (NC - len(t['classes'])) + [t.get('window', ''), adj_text(t.get('adjust'))])
        for t in e.get('matrix_tools', []):
            R, C = t['rows'], t['cols']
            mt.append([k, t.get('order', ''), t['label'], INDEX_NAME[R['idx']], R['fm']] + R['breaks'] + [None] * (NB - len(R['breaks'])) +
                      [INDEX_NAME[C['idx']], C['fm']] + C['breaks'] + [None] * (NB - len(C['breaks'])) +
                      [' '.join(map(str, r)) for r in t['matrix']] + [None] * (NC - len(t['matrix'])) +
                      t['classes'] + [None] * (NC - len(t['classes'])) + [t.get('window', ''), adj_text(t.get('adjust'))])
    for s in (info, tl, mt):   # text cells stay text in Excel (no 05-01 -> 1-May, no 44302, 45220 -> number)
        for row in s.iter_rows(min_row=2):
            for c in row:
                h = s.cell(1, c.column).value
                if not (h.startswith(('Break', 'Row break', 'Column break')) or h in ('Order', 'Calibration from', 'Calibration to', 'Large fire (acres)', 'Multi-fire day (fires)')):
                    c.number_format = '@'
                if h == 'Notes':
                    c.alignment = Alignment(wrap_text=False)
    # dropdowns and number checks (Excel stops most typos before they reach GitHub)
    lists = wb.create_sheet('Lists')
    for i, k in enumerate(ids, 1): lists.cell(i, 1, k)
    for i, k in enumerate(INDEX, 1): lists.cell(i, 2, k)
    for i, k in enumerate(FUEL, 1): lists.cell(i, 3, k)
    lists.sheet_state = 'hidden'
    n = 5000

    def dv(s, rng, formula, msg):
        v = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True, errorTitle='Not on the list', error=msg)
        s.add_data_validation(v); v.add(rng)
    dv(tl, f'A2:A{n}', f'=Lists!$A$1:$A${len(ids)}', 'Use an FDRA ID from the FDOP info sheet.')
    dv(info, f'A2:A{n}', f'=Lists!$A$1:$A${len(ids)}', 'Use an FDRA ID from the map. A new FDRA needs its shape added first.')
    dv(tl, f'D2:D{n}', f'=Lists!$B$1:$B${len(INDEX)}', 'Pick the index from the list.')
    dv(info, f'E2:E{n}', f'=Lists!$B$1:$B${len(INDEX)}', 'Pick the index from the list.')
    dv(tl, f'E2:E{n}', f'=Lists!$C$1:$C${len(FUEL)}', 'Fuel model V, W, X, Y or Z.')
    dv(info, f'F2:F{n}', f'=Lists!$C$1:$C${len(FUEL)}', 'Fuel model V, W, X, Y or Z.')
    from openpyxl.utils import get_column_letter as L_
    dv(mt, f'A2:A{n}', f'=Lists!$A$1:$A${len(ids)}', 'Use an FDRA ID from the FDOP info sheet.')
    ci = {h: i + 1 for i, h in enumerate(MAT_COLS)}
    for h in ('Row index', 'Column index'):
        dv(mt, f'{L_(ci[h])}2:{L_(ci[h])}{n}', f'=Lists!$B$1:$B${len(INDEX)}', 'Pick the index from the list.')
    for h in ('Row fuel model', 'Column fuel model'):
        dv(mt, f'{L_(ci[h])}2:{L_(ci[h])}{n}', f'=Lists!$C$1:$C${len(FUEL)}', 'Fuel model V, W, X, Y or Z.')
    for sh, rngs in ((tl, [f'F2:{L_(5 + NB)}{n}']),
                     (mt, [f'{L_(ci["Row break 1"])}2:{L_(ci["Row break " + str(NB)])}{n}',
                           f'{L_(ci["Column break 1"])}2:{L_(ci["Column break " + str(NB)])}{n}'])):
        num = DataValidation(type='decimal', operator='between', formula1='-1000', formula2='100000', allow_blank=True,
                             showErrorMessage=True, errorTitle='Number only', error='Breakpoints are numbers (for example 42 or 55.9).')
        sh.add_data_validation(num)
        for r_ in rngs: num.add(r_)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------- read + check
def _s(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()


def _mmdd(v, err):
    if v in (None, ''): return None
    if isinstance(v, (dt.date, dt.datetime)): return v.strftime('%m-%d')
    s = _s(v)
    m = re.fullmatch(r'(\d{1,2})[-/](\d{1,2})', s)
    if not m: err(f'season date "{s}" should look like 05-01'); return None
    mo, d = int(m.group(1)), int(m.group(2))
    if not (1 <= mo <= 12 and 1 <= d <= 31): err(f'season date "{s}" is not a real month-day'); return None
    return f'{mo:02d}-{d:02d}'


def _int(v, what, err, lo=None, hi=None):
    if v in (None, ''): return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        err(f'{what} "{v}" is not a number'); return None
    if not f.is_integer(): err(f'{what} "{v}" should be a whole number'); return None
    f = int(f)
    if (lo is not None and f < lo) or (hi is not None and f > hi): err(f'{what} {f} is out of range'); return None
    return f


def _num(v):
    f = float(v)
    return int(f) if f.is_integer() else round(f, 3)


def adj_text(adj):
    return '; '.join(f"{ADJ_NAME[a['var']]} {a['op']} {a['value']:g} : {a['adjust']:+d}" + (f" @{a['window']}" if a.get('window') else '')
                     for a in adj or [])


def _window(v, err):
    w = _s(v).upper().replace(' ', '')
    if not w:
        return ''
    if not re.fullmatch(r'(O1|F[1-7])(\+(O1|F[1-7]))*', w):
        err(f'days used "{_s(v)}" should look like F1 or O1+F1+F2+F3 (O1 = observed, F1-F7 = forecast days)'); return None
    return w


def _adjust(v, err):
    out = []
    for part in [p.strip() for p in re.split(r'[;\n]+', _s(v)) if p.strip()]:
        m = re.fullmatch(r'(.+?)\s*(<=|>=|<|>)\s*(-?\d+(?:\.\d+)?)\s*:\s*([+-]?\d)\s*(?:@\s*(\S+))?', part)
        var = m and ADJ_VARS.get(re.sub(r'\s+', ' ', m.group(1)).upper().replace('100 HR', '100-HR').replace('1000 HR', '1000-HR'))
        if not m or not var:
            err(f'adjustment "{part}" should look like  Tmax < 57 : -1  (variables: {", ".join(ADJ_NAME.values())})'); return None
        a = {'var': var, 'op': m.group(2), 'value': _num(m.group(3)), 'adjust': int(m.group(4))}
        if a['adjust'] == 0 or abs(a['adjust']) > 2:
            err(f'adjustment "{part}": the change should be -1, +1, -2 or +2'); return None
        if m.group(5):
            w = _window(m.group(5), err)
            if w is None: return None
            a['window'] = w
        out.append(a)
    return out


def read(path: Path = BOOK):
    """-> (start dict, problems [ (sheet, row, fdra_id, message) ], bad fdra ids)"""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    probs, bad = [], set()
    ids = {f['properties']['fdra_id'] for f in json.loads(FDRAS.read_text())['features']}
    for name, cols in (('FDOP info', INFO_COLS), ('Thresholds', TOOL_COLS)):
        if name not in wb.sheetnames:
            return None, [(name, 0, '', f'sheet "{name}" is missing')], set()
        got = [_s(c.value) for c in wb[name][1]][:len(cols)]
        miss = [c for c in cols if c not in got]
        if miss:
            return None, [(name, 1, '', f'missing column(s): {", ".join(miss)} (keep the headings as they are)')], set()

    def rows(name):
        s = wb[name]; hdr = [_s(c.value) for c in s[1]]
        for r in s.iter_rows(min_row=2):
            v = {h: c.value for h, c in zip(hdr, r) if h}
            if any(_s(x) for x in v.values()):
                yield r[0].row, v

    start = {}
    for rn, v in rows('FDOP info'):
        k = _s(v['FDRA ID'])
        def err(m, k=k, rn=rn): probs.append(('FDOP info', rn, k, m)); bad.add(k)
        if not k: err('FDRA ID is blank'); continue
        if k not in ids: err(f'FDRA ID "{k}" is not on the map (check spelling, or add the shape first)'); continue
        if k in start: err('FDRA ID is listed twice'); continue
        e = {'source': _s(v['FDOP source'])}
        st = re.split(r'[,;\s]+', _s(v['SIG stations'])) if _s(v['SIG stations']) else []
        if any(not s.isdigit() for s in st): err(f'SIG stations "{_s(v["SIG stations"])}" should be station numbers separated by commas')
        e['stations'] = [int(s) for s in st if s.isdigit()]
        fm = _s(v['Fuel model']).upper()
        if fm and fm not in FUEL: err(f'fuel model "{fm}" should be V, W, X, Y or Z')
        e['fm'] = fm
        ix = _s(v['Primary index'])
        if ix and ix not in INDEX: err(f'primary index "{ix}" is not on the list')
        e['idx'] = INDEX.get(ix, '')
        w = _s(v['Station weights'])
        if w:
            try:
                e['weights'] = {a.strip(): _num(b) for a, b in (p.split(':') for p in re.split(r'[,;]+', w) if p.strip())}
            except ValueError:
                err(f'station weights "{w}" should look like 44503:30, 44203:30')
        a, b = _mmdd(v['Season start (MM-DD)'], err), _mmdd(v['Season end (MM-DD)'], err)
        if a and b: e['season'] = [a, b]
        elif a or b: err('give both season start and season end, or neither')
        y = dt.date.today().year
        c1, c2 = _int(v['Calibration from'], 'calibration from', err, 1950, y), _int(v['Calibration to'], 'calibration to', err, 1950, y)
        if c1 and c2:
            if c2 < c1: err('calibration to is before calibration from')
            e['calib'] = [c1, c2]
        elif c1 or c2: err('give both calibration years, or neither')
        lg, mf = _int(v['Large fire (acres)'], 'large fire acres', err, 1), _int(v['Multi-fire day (fires)'], 'multi-fire day', err, 2)
        if lg: e['large'] = lg
        if mf: e['mfd'] = mf
        cl = [c.strip() for c in _s(v['Class names (primary)']).split('|') if c.strip()]
        if cl: e['classes'] = cl
        adv = _s(v['Advanced (JSON)'])
        if adv:
            try:
                x = json.loads(adv); assert isinstance(x, dict); e.update(x)
            except Exception:
                err('Advanced (JSON) is not valid JSON - leave it blank unless you know it is needed')
        e['notes'] = _s(v['Notes'])
        e['_row'] = rn
        start[k] = e
    tools = {}
    for rn, v in rows('Thresholds'):
        k = _s(v['FDRA ID'])
        def err(m, k=k, rn=rn): probs.append(('Thresholds', rn, k, m)); bad.add(k)
        if not k: err('FDRA ID is blank'); continue
        if k not in start:
            err(f'FDRA ID "{k}" has no row on the FDOP info sheet' if k in ids else f'FDRA ID "{k}" is not on the map'); continue
        ix = _s(v['Index'])
        if ix not in INDEX: err(f'index "{ix}" is not on the list'); continue
        fm = _s(v['Fuel model']).upper()
        if fm not in FUEL: err(f'fuel model "{fm}" should be V, W, X, Y or Z'); continue
        br, ok = [], True
        cells = [v[f'Break {i}'] for i in range(1, NB + 1)]
        while cells and cells[-1] in (None, ''): cells.pop()
        for i, c in enumerate(cells, 1):
            try:
                br.append(_num(c))
            except (TypeError, ValueError):
                err(f'Break {i} "{c}" is not a number (blank in the middle, or text)'); ok = False
        if not ok: continue
        if not br: err('no breakpoints'); continue
        if any(b2 <= b1 for b1, b2 in zip(br, br[1:])): err(f'breakpoints {br} must go up from left to right'); continue
        cl = [_s(v[f'Class {i}']) for i in range(1, NC + 1)]
        while cl and not cl[-1]: cl.pop()
        if any(not c for c in cl): err('a class name in the middle is blank'); continue
        if len(cl) != len(br) + 1:
            err(f'{len(br)} breakpoints need {len(br) + 1} class names (found {len(cl)})'); continue
        order = _int(v['Order'], 'order', err)
        win, adj = _window(v.get('Days used'), err), _adjust(v.get('Adjustments'), err)
        if win is None or adj is None: continue
        t = {'label': _s(v['Tool name']) or f'{ix}-{fm}', 'idx': INDEX[ix], 'fm': fm, 'breaks': br, 'classes': cl}
        if win: t['window'] = win
        if adj: t['adjust'] = adj
        tools.setdefault(k, []).append((order if order is not None else 999, rn, t))
    mats = {}
    if 'Matrix tools' in wb.sheetnames:
        got = [_s(c.value) for c in wb['Matrix tools'][1]]
        miss = [c for c in MAT_COLS if c not in got and c not in TOOL_OPT]
        if miss:
            return None, [('Matrix tools', 1, '', f'missing column(s): {", ".join(miss)} (keep the headings as they are)')], set()
        for rn, v in rows('Matrix tools'):
            k = _s(v['FDRA ID'])
            def err(m, k=k, rn=rn): probs.append(('Matrix tools', rn, k, m)); bad.add(k)
            if not k: err('FDRA ID is blank'); continue
            if k not in start:
                err(f'FDRA ID "{k}" has no row on the FDOP info sheet' if k in ids else f'FDRA ID "{k}" is not on the map'); continue
            ax, ok = {}, True
            for side, lab in (('rows', 'Row'), ('cols', 'Column')):
                ix, fm = _s(v[f'{lab} index']), _s(v[f'{lab} fuel model']).upper()
                if ix not in INDEX: err(f'{lab.lower()} index "{ix}" is not on the list'); ok = False; break
                if fm not in FUEL: err(f'{lab.lower()} fuel model "{fm}" should be V, W, X, Y or Z'); ok = False; break
                cells = [v[f'{lab} break {i}'] for i in range(1, NB + 1)]
                while cells and cells[-1] in (None, ''): cells.pop()
                try:
                    br = [_num(c) for c in cells]
                except (TypeError, ValueError):
                    err(f'{lab.lower()} breakpoints {cells} must all be numbers'); ok = False; break
                if not br or any(b2 <= b1 for b1, b2 in zip(br, br[1:])):
                    err(f'{lab.lower()} breakpoints {br} must be given and go up from left to right'); ok = False; break
                ax[side] = {'idx': INDEX[ix], 'fm': fm, 'breaks': br}
            if not ok: continue
            cl = [_s(v[f'Class {i}']) for i in range(1, NC + 1)]
            while cl and not cl[-1]: cl.pop()
            if len(cl) < 2 or any(not c for c in cl): err('give at least two class names, with no blanks in the middle'); continue
            nr, nc_ = len(ax['rows']['breaks']) + 1, len(ax['cols']['breaks']) + 1
            mat = []
            for i in range(1, nr + 1):
                txt = _s(v.get(f'Row {i} levels'))
                try:
                    lv = [int(x) for x in re.split(r'[\s,]+', txt) if x]
                except ValueError:
                    err(f'Row {i} levels "{txt}" should be whole numbers separated by spaces'); ok = False; break
                if len(lv) != nc_:
                    err(f'Row {i} levels needs {nc_} numbers, one per column class (found {len(lv)})'); ok = False; break
                if any(x < 1 or x > len(cl) for x in lv):
                    err(f'Row {i} levels must be between 1 and {len(cl)} (the number of classes)'); ok = False; break
                mat.append(lv)
            if not ok: continue
            extra_rows = [i for i in range(nr + 1, NC + 1) if _s(v.get(f'Row {i} levels'))]
            if extra_rows: err(f'{nr - 1} row breakpoints give {nr} rows, but Row {extra_rows[0]} levels is filled in'); continue
            win, adj = _window(v.get('Days used'), err), _adjust(v.get('Adjustments'), err)
            if win is None or adj is None: continue
            order = _int(v['Order'], 'order', err)
            t = {'type': 'matrix', 'label': _s(v['Tool name']) or 'Matrix', **ax, 'matrix': mat, 'classes': cl}
            if win: t['window'] = win
            if adj: t['adjust'] = adj
            mats.setdefault(k, []).append((order if order is not None else 999, rn, t))
    for k, e in start.items():
        allt = sorted([(o, 0, rn, t) for o, rn, t in tools.get(k, [])] + [(o, 1, rn, t) for o, rn, t in mats.get(k, [])],
                      key=lambda x: (x[0], x[1], x[2]))
        for i, (_, _, _, t) in enumerate(allt, 1):
            t['order'] = i
        ts = [t for _, kind, _, t in allt if kind == 0]
        if mats.get(k):
            e['matrix_tools'] = [t for _, kind, _, t in allt if kind == 1]
        if ts:
            e['tools'] = ts
            e['breaks'] = {}
            for t in ts: e['breaks'].setdefault(t['idx'], t['breaks'])
            if not e.get('idx'): e['idx'] = ts[0]['idx']
            if not e.get('fm'): e['fm'] = ts[0]['fm']
            if 'classes' in e and e['idx'] in e['breaks'] and len(e['classes']) != len(e['breaks'][e['idx']]) + 1:
                probs.append(('FDOP info', e['_row'], k, f'{len(e["classes"])} primary class names but the first {INDEX_NAME[e["idx"]]} '
                              f'tool has {len(e["breaks"][e["idx"]])} breakpoints')); bad.add(k)
            if 'classes' not in e:
                first = next((t for t in ts if t['idx'] == e['idx']), ts[0]); e['classes'] = first['classes']
        elif e.get('idx') and not e.get('source'):
            pass
    keep = {}
    for k, e in start.items():
        e.pop('_row', None)
        if not e.get('source') and not e.get('tools') and not e.get('matrix_tools'):
            continue
        F_ = ('source', 'stations', 'fm', 'idx', 'season', 'calib', 'breaks', 'classes', 'large', 'mfd', 'tools', 'matrix_tools')
        keep[k] = {x: e[x] for x in F_ + ('notes',) if x in e and e[x] not in ('', None)}
        keep[k].update({x: e[x] for x in e if x not in keep[k] and x != 'notes' and x not in F_})
        keep[k].setdefault('notes', e.get('notes', ''))
    return keep, probs, bad


def build(out: Path, prev: Path | None = None):
    """Write out/fdop_start.json from the workbook; FDRAs with problems keep their previous values."""
    out.mkdir(parents=True, exist_ok=True)
    prev_d = {}
    for p in (prev, out / 'fdop_start.json', LEGACY):
        if p and Path(p).exists():
            try:
                prev_d = json.loads(Path(p).read_text()); break
            except ValueError:
                pass
    if not BOOK.exists():
        d, probs, bad = {k: v for k, v in prev_d.items() if k != '_about'}, [('', 0, '', 'workbook not found; previous thresholds kept')], set()
    else:
        d, probs, bad = read()
        if d is None:   # structure broken: keep everything as it was
            d, bad = {k: v for k, v in prev_d.items() if k != '_about'}, set()
        else:
            for k in bad:
                if k in prev_d: d[k] = prev_d[k]
                else: d.pop(k, None)
    res = {'_about': ABOUT, **dict(sorted(d.items()))}
    (out / 'fdop_start.json').write_text(json.dumps(res, indent=1, ensure_ascii=False))
    (out / 'thresholds_check.json').write_text(json.dumps({
        'generated': dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%d %H:%M UTC'), 'fdras': len(d),
        'problems': [{'sheet': s, 'row': r, 'fdra_id': k, 'problem': m} for s, r, k, m in probs],
        'kept_previous': sorted(bad)}, indent=1))
    return res, probs


def report(probs):
    if not probs:
        return '### Thresholds check: no problems found\nThe site uses the new thresholds at the next build.'
    lines = ['### Thresholds check: problems found',
             'These FDRAs keep their last good thresholds until the rows are fixed and the file is uploaded again.', '',
             '| Sheet | Row | FDRA | Problem |', '|---|---|---|---|']
    lines += [f'| {s} | {r} | {k} | {m} |' for s, r, k, m in probs]
    return '\n'.join(lines)


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'check'
    if cmd == 'export':
        src = Path(sys.argv[2]) if len(sys.argv) > 2 else LEGACY
        d = json.loads(src.read_text()); d.pop('_about', None)
        export(d); print('wrote', BOOK)
    elif cmd == 'build':
        _, probs = build(Path(sys.argv[2]), Path(sys.argv[3]) if len(sys.argv) > 3 else None)
        print(report(probs))
    else:
        d, probs, bad = read()
        print(report(probs))
        if d is not None: print(f'\n{len(d)} FDRAs with FDOP thresholds read.')
        sys.exit(1 if probs else 0)
